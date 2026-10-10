from __future__ import annotations

import json
import subprocess
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

from conftest import SKILL_DIRS, SKILL_IDS, config_data, load_engine, project_data, template_path


def _doc_all_text(doc) -> str:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    for section in doc.sections:
        texts.extend(p.text for p in section.header.paragraphs)
        texts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(texts)


@pytest.mark.parametrize("skill_id", SKILL_IDS)
def test_end_to_end_fill(skill_id, tmp_path):
    eng = load_engine(skill_id)
    config = config_data(skill_id)
    project = project_data(skill_id)
    plan = eng.loader.build_plan(config, project)

    doc = Document(str(template_path(skill_id)))
    eng.docx_ops.apply_plan(doc, plan)
    output = tmp_path / f"filled-{skill_id}.docx"
    doc.save(str(output))

    result = Document(str(output))
    text = _doc_all_text(result)
    assert config["project"]["name"]["content"] in text
    assert config["document"]["id"]["content"] in text
    assert "外部型号+产品名称" not in text
    assert "XXXX年XX月XX日" not in text
    assert "产品型号-XXXX" not in text

    update_fields = result.settings.element.find(qn("w:updateFields"))
    assert update_fields is not None
    assert update_fields.get(qn("w:val")) == "true"


def _run_cli(skill_id, config_path, project_path, output_path) -> subprocess.CompletedProcess:
    scripts_dir = SKILL_DIRS[skill_id] / "scripts"
    return subprocess.run(
        [
            sys.executable,
            str(scripts_dir / "main.py"),
            "--template",
            str(template_path(skill_id)),
            "--config",
            str(config_path),
            "--project",
            str(project_path),
            "--output",
            str(output_path),
        ],
        capture_output=True,
        text=True,
    )


@pytest.mark.parametrize("skill_id", SKILL_IDS)
def test_cli_success(skill_id, tmp_path):
    output = tmp_path / f"cli-{skill_id}.docx"
    proc = _run_cli(
        skill_id,
        SKILL_DIRS[skill_id] / "templates" / "config.json",
        SKILL_DIRS[skill_id] / "templates" / "project.json",
        output,
    )
    assert proc.returncode == 0, proc.stderr
    assert output.exists()
    assert "文档已生成" in proc.stdout


@pytest.mark.parametrize("skill_id", SKILL_IDS)
def test_cli_data_error_exit_code(skill_id, tmp_path):
    broken_config = config_data(skill_id)
    del broken_config["project"]["name"]
    broken_path = tmp_path / "broken-config.json"
    broken_path.write_text(json.dumps(broken_config, ensure_ascii=False), encoding="utf-8")

    proc = _run_cli(
        skill_id,
        broken_path,
        SKILL_DIRS[skill_id] / "templates" / "project.json",
        tmp_path / "should-not-exist.docx",
    )
    assert proc.returncode == 1
    assert "数据校验失败" in proc.stderr
    assert "Traceback" not in proc.stderr
