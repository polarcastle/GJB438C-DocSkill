from __future__ import annotations

import pytest
from docx import Document

from conftest import SKILL_IDS, load_engine


@pytest.mark.parametrize("skill_id", SKILL_IDS)
class TestSetParagraphText:
    def test_replaces_text_and_keeps_style_run(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        paragraph = doc.add_paragraph()
        style_run = paragraph.add_run("旧")
        style_run.bold = True
        paragraph.add_run("文本")
        eng.docx_ops.set_paragraph_text(paragraph, "新文本")
        assert paragraph.text == "新文本"
        assert paragraph.runs[0].bold
        assert all(run.text == "" for run in paragraph.runs[1:])

    def test_adds_run_when_paragraph_empty(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        paragraph = doc.add_paragraph()
        eng.docx_ops.set_paragraph_text(paragraph, "内容")
        assert paragraph.text == "内容"


@pytest.mark.parametrize("skill_id", SKILL_IDS)
class TestSetParagraphsFromText:
    def test_multiline_clones_paragraphs(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        paragraph = doc.add_paragraph("占位")
        eng.docx_ops.set_paragraphs_from_text(paragraph, "第一行\n\n第二行\r\n第三行")
        assert [p.text for p in doc.paragraphs] == ["第一行", "第二行", "第三行"]

    def test_blank_text_rejected(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        paragraph = doc.add_paragraph("占位")
        with pytest.raises(eng.errors.StrictTemplateError, match="不能为空"):
            eng.docx_ops.set_paragraphs_from_text(paragraph, "  \n  ")


@pytest.mark.parametrize("skill_id", SKILL_IDS)
class TestReplacements:
    def test_exact_replacement_count_mismatch(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        doc.add_paragraph("目标")
        doc.add_paragraph("目标")
        with pytest.raises(eng.errors.StrictTemplateError, match="期望命中 1 个"):
            eng.docx_ops.apply_exact_replacement(doc, "目标", "新", 1, "k")

    def test_exact_replacement_reaches_table_cells(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "目标"
        eng.docx_ops.apply_exact_replacement(doc, "目标", "新", 1, "k")
        assert table.cell(0, 0).text == "新"

    def test_paragraph_replacement_multiline(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        doc.add_paragraph("前文")
        doc.add_paragraph("这里是锚点说明文字")
        doc.add_paragraph("后文")
        eng.docx_ops.apply_paragraph_replacement(doc, "锚点说明", "第一\n第二", 1, "k")
        assert [p.text for p in doc.paragraphs] == ["前文", "第一", "第二", "后文"]


@pytest.mark.parametrize("skill_id", SKILL_IDS)
class TestCaptionNumbering:
    def _doc_with_captions(self) -> Document:
        doc = Document()
        doc.add_heading("第一章", level=1)
        doc.add_paragraph("表X-X 能力表")
        doc.add_paragraph("图X-X 结构图")
        doc.add_heading("第二章", level=1)
        doc.add_paragraph("表X-X 接口表")
        return doc

    def test_numbers_per_chapter_and_kind(self, skill_id):
        eng = load_engine(skill_id)
        doc = self._doc_with_captions()
        eng.docx_ops.apply_caption_numbering(doc)
        texts = [p.text for p in doc.paragraphs]
        assert "表1-1 能力表" in texts
        assert "图1-1 结构图" in texts
        assert "表2-1 接口表" in texts

    def test_caption_before_any_heading_rejected(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        doc.add_paragraph("表X-X 孤儿表")
        with pytest.raises(eng.errors.StrictTemplateError, match="任何一级章节之前"):
            eng.docx_ops.apply_caption_numbering(doc)

    def test_explicit_wrong_chapter_rejected(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        doc.add_heading("第一章", level=1)
        doc.add_paragraph("表2-X 错位表")
        with pytest.raises(eng.errors.StrictTemplateError, match="章号与当前位置不一致"):
            eng.docx_ops.apply_caption_numbering(doc)


@pytest.mark.parametrize("skill_id", SKILL_IDS)
class TestTableFill:
    def _doc_with_table(self) -> Document:
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "名称"
        table.cell(0, 1).text = "说明"
        table.cell(1, 0).text = "占位一"
        table.cell(2, 0).text = "占位二"
        return doc

    def _rule(self, eng, rows):
        return eng.models.TableFillRule(
            key="3.T1",
            locator_rows=(("名称", "说明"),),
            columns=("c1", "c2"),
            rows=rows,
            data_start_row=1,
            preserve_tail_rows=0,
        )

    def test_fills_and_extends_rows(self, skill_id):
        eng = load_engine(skill_id)
        doc = self._doc_with_table()
        rows = tuple((f"r{i}c1", f"r{i}c2") for i in range(4))
        eng.docx_ops.apply_table_fill(doc, self._rule(eng, rows))
        table = doc.tables[0]
        assert len(table.rows) == 5
        assert table.cell(1, 0).text == "r0c1"
        assert table.cell(4, 1).text == "r3c2"

    def test_leftover_template_rows_cleared(self, skill_id):
        eng = load_engine(skill_id)
        doc = self._doc_with_table()
        eng.docx_ops.apply_table_fill(doc, self._rule(eng, (("仅一行", "x"),)))
        table = doc.tables[0]
        assert table.cell(1, 0).text == "仅一行"
        assert table.cell(2, 0).text == ""
        assert table.cell(2, 1).text == ""

    def test_locator_mismatch_rejected(self, skill_id):
        eng = load_engine(skill_id)
        doc = self._doc_with_table()
        rule = eng.models.TableFillRule(
            key="3.T9",
            locator_rows=(("不存在", "的表头"),),
            columns=("c1", "c2"),
            rows=(("a", "b"),),
            data_start_row=1,
            preserve_tail_rows=0,
        )
        with pytest.raises(eng.errors.StrictTemplateError, match="期望命中 1 张表"):
            eng.docx_ops.apply_table_fill(doc, rule)


@pytest.mark.parametrize("skill_id", SKILL_IDS)
class TestHelpers:
    def test_heading_separator(self, skill_id):
        eng = load_engine(skill_id)
        assert eng.docx_ops._heading_separator("3.2 软件能力需求") == " "
        assert eng.docx_ops._heading_separator("3.2软件能力需求") == ""
        assert eng.docx_ops._heading_separator("") == ""

    def test_find_single_body_paragraph_requires_unique(self, skill_id):
        eng = load_engine(skill_id)
        doc = Document()
        doc.add_paragraph("重复锚点")
        doc.add_paragraph("重复锚点")
        with pytest.raises(eng.errors.StrictTemplateError, match="应为 1 个"):
            eng.docx_ops.find_single_body_paragraph(doc, "重复锚点")
